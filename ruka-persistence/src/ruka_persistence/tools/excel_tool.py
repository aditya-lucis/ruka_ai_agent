from pathlib import Path
import openpyxl

class ExcelTool:
    def read_tables(self, path: str | Path) -> dict:
        wb = openpyxl.load_workbook(path, data_only=True)
        result = {}
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            # Read first few rows as a table for simulation
            data = []
            for row in ws.iter_rows(values_only=True, max_row=10):
                if any(row):
                    data.append(row)
            result[sheet] = data
        return result
